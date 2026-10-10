"""Independent, read-only comparison of the original IVG workbook.

The workbook is an evidence source, not a training source.  This module
recomputes its numeric formulas from the stored inputs and compares a
stratified sample of the original synthetic inputs with the separate v3
3-uF detailed RLC model under both declared graph assumptions.  Workbook
formula results, model-theoretical values, and any measured observations are
kept in separate fields in the report.

``openpyxl`` is used when available.  A small XML reader is included for the
portable runtime so the comparison remains useful when the optional workbook
package is unavailable.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import math
import os
from pathlib import Path
import random
import re
import statistics
from typing import Any, Iterable, Mapping, Sequence
import zipfile
import xml.etree.ElementTree as ET


WORKBOOK_NAME = "Hybrid_Physics_ML_Impulse_Generator_Optimiser.xlsx"
EXPECTED_SOURCE_SHA256 = "e855eadbe4d6da2579316e89e025e6990f9ba3c522a5af96d1dd2a5522d77480"
COMPARISON_SCHEMA_VERSION = "powernext_v3_excel_comparison_1.0"


def _finite(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    return value if math.isfinite(value) else None


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def locate_workbook(root: str | Path | None = None) -> Path:
    """Locate the preserved original workbook without editing it."""

    base = Path(root) if root is not None else Path(__file__).resolve().parent
    candidates = [
        base / "evidence" / "sources" / WORKBOOK_NAME,
        base / "sources" / "original" / WORKBOOK_NAME,
        base / "competition-materials" / WORKBOOK_NAME,
        base / WORKBOOK_NAME,
    ]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    # Useful when called from a subdirectory in a checkout.
    for parent in [base, *base.parents]:
        for relative in (Path("evidence/sources") / WORKBOOK_NAME, Path("sources/original") / WORKBOOK_NAME, Path("competition-materials") / WORKBOOK_NAME):
            candidate = parent / relative
            if candidate.is_file():
                return candidate
    raise FileNotFoundError(f"Unable to locate preserved {WORKBOOK_NAME}")


def _column_index(reference: str) -> int:
    letters = re.match(r"([A-Z]+)", reference.upper())
    if not letters:
        raise ValueError(reference)
    value = 0
    for letter in letters.group(1):
        value = value * 26 + ord(letter) - ord("A") + 1
    return value


def _cell_parts(reference: str) -> tuple[int, int]:
    match = re.fullmatch(r"([A-Z]+)([0-9]+)", reference.upper())
    if not match:
        raise ValueError(reference)
    return int(match.group(2)), _column_index(match.group(1))


def _openpyxl_cells(path: Path) -> tuple[dict[str, dict[str, dict[str, Any]]], dict[str, dict[str, dict[str, Any]]], str]:
    from openpyxl import load_workbook

    formulas = load_workbook(path, data_only=False, read_only=False)
    cached = load_workbook(path, data_only=True, read_only=False)
    formula_sheets: dict[str, dict[str, dict[str, Any]]] = {}
    cached_sheets: dict[str, dict[str, dict[str, Any]]] = {}
    for sheet in formulas.worksheets:
        formula_cells: dict[str, dict[str, Any]] = {}
        cached_cells: dict[str, dict[str, Any]] = {}
        cached_sheet = cached[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                formula = cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None
                formula_cells[cell.coordinate] = {"value": cell.value, "formula": formula}
                cached_cell = cached_sheet[cell.coordinate]
                cached_cells[cell.coordinate] = {"value": cached_cell.value, "formula": formula}
        formula_sheets[sheet.title] = formula_cells
        cached_sheets[sheet.title] = cached_cells
    return formula_sheets, cached_sheets, "openpyxl"


_NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _xml_text(element: ET.Element | None) -> str | None:
    if element is None:
        return None
    return "".join(element.itertext())


def _xml_value(raw: str | None, kind: str | None, shared_strings: list[str]) -> Any:
    if raw is None:
        return None
    if kind == "s":
        try:
            return shared_strings[int(raw)]
        except (IndexError, ValueError):
            return raw
    if kind == "b":
        return raw == "1"
    if kind == "str":
        return raw
    try:
        number = float(raw)
        return int(number) if number.is_integer() else number
    except ValueError:
        return raw


def _xml_cells(path: Path) -> tuple[dict[str, dict[str, dict[str, Any]]], dict[str, dict[str, dict[str, Any]]], str]:
    """Read cell/formula/cached-value triples directly from XLSX XML."""

    with zipfile.ZipFile(path) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in root.findall(f"{_NS_MAIN}si"):
                shared_strings.append(_xml_text(item) or "")
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_targets: dict[str, str] = {}
        for relation in relations:
            rid = relation.attrib.get("Id")
            target = relation.attrib.get("Target", "")
            if rid:
                clean_target = target.lstrip("/")
                rel_targets[rid] = clean_target if clean_target.startswith("xl/") else "xl/" + clean_target
        formulas_by_sheet: dict[str, dict[str, dict[str, Any]]] = {}
        cached_by_sheet: dict[str, dict[str, dict[str, Any]]] = {}
        for sheet in workbook.findall(f"{_NS_MAIN}sheets/{_NS_MAIN}sheet"):
            title = sheet.attrib.get("name", "Sheet")
            rid = sheet.attrib.get(f"{_NS_REL}id")
            target = rel_targets.get(rid or "")
            if not target or target not in archive.namelist():
                continue
            sheet_root = ET.fromstring(archive.read(target))
            formula_cells: dict[str, dict[str, Any]] = {}
            cached_cells: dict[str, dict[str, Any]] = {}
            for cell in sheet_root.findall(f".//{_NS_MAIN}c"):
                reference = cell.attrib.get("r")
                if not reference:
                    continue
                kind = cell.attrib.get("t")
                value = _xml_value(_xml_text(cell.find(f"{_NS_MAIN}v")), kind, shared_strings)
                formula_element = cell.find(f"{_NS_MAIN}f")
                formula_text = _xml_text(formula_element)
                formula = ("=" + formula_text) if formula_text is not None else None
                if formula is not None:
                    formula_cells[reference] = {"value": formula, "formula": formula}
                    cached_cells[reference] = {"value": value, "formula": formula}
                elif value is not None:
                    formula_cells[reference] = {"value": value, "formula": None}
                    cached_cells[reference] = {"value": value, "formula": None}
            formulas_by_sheet[title] = formula_cells
            cached_by_sheet[title] = cached_cells
    return formulas_by_sheet, cached_by_sheet, "zipxml"


def read_workbook(path: str | Path) -> dict[str, Any]:
    """Read workbook values and formulas through the available backend."""

    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    try:
        formulas, cached, backend = _openpyxl_cells(source)
    except (ImportError, ModuleNotFoundError):
        formulas, cached, backend = _xml_cells(source)
    return {"path": str(source), "sha256": _sha256(source), "backend": backend, "formulas": formulas, "cached": cached}


def _value(book: Mapping[str, Any], sheet: str, cell: str, *, cached: bool = True) -> Any:
    table = book["cached" if cached else "formulas"].get(sheet, {})
    return table.get(cell, {}).get("value")


def _sheet_rows(book: Mapping[str, Any], sheet: str) -> list[dict[str, Any]]:
    cells = book["cached"].get(sheet, {})
    rows: dict[int, dict[str, Any]] = defaultdict(dict)
    for coordinate, item in cells.items():
        try:
            row, col = _cell_parts(coordinate)
        except ValueError:
            continue
        # Convert column number back to a compact key for predictable records.
        number = col
        letters = ""
        while number:
            number, remainder = divmod(number - 1, 26)
            letters = chr(65 + remainder) + letters
        rows[row][letters] = item.get("value")
    return [rows[index] for index in sorted(rows)]


def _excel_round(value: float, digits: int = 0) -> float:
    multiplier = 10.0 ** digits
    return math.floor(value * multiplier + 0.5) / multiplier if value >= 0 else math.ceil(value * multiplier - 0.5) / multiplier


def _calculator_values(book: Mapping[str, Any], overrides: Mapping[str, Any] | None = None) -> dict[str, Any]:
    """Recompute the original calculator formulas from visible inputs."""

    values = {
        "max_system_kV": _value(book, "Hybrid Calculator", "B5"),
        "impulse_type": _value(book, "Hybrid Calculator", "B6"),
        "test_kV": _value(book, "Hybrid Calculator", "B7"),
        "load_pF": _value(book, "Hybrid Calculator", "B8"),
        "divider_pF": _value(book, "Hybrid Calculator", "B9"),
        "stray_pF": _value(book, "Hybrid Calculator", "B10"),
        "stray_inductance_uH": _value(book, "Hybrid Calculator", "B11"),
        "efficiency": _value(book, "Hybrid Calculator", "B12"),
        "max_stages": _value(book, "Hybrid Calculator", "B13"),
        "max_stage_voltage_kV": _value(book, "Hybrid Calculator", "B14"),
        "stage_capacitance_uF": _value(book, "Hybrid Calculator", "B15"),
    }
    if overrides:
        values.update(overrides)
    impulse = str(values["impulse_type"])
    target_front = 1.2 if impulse == "Lightning" else 250.0
    target_tail = 50.0 if impulse == "Lightning" else 2_500.0
    load_total = sum(float(values[name]) for name in ("load_pF", "divider_pF", "stray_pF"))
    stages = math.ceil(float(values["test_kV"]) / (float(values["max_stage_voltage_kV"]) * float(values["efficiency"])))
    charge_stage = float(values["test_kV"]) / (float(values["efficiency"]) * stages)
    utilization = charge_stage / float(values["max_stage_voltage_kV"])
    generator_cap = float(values["stage_capacitance_uF"]) / stages
    total_front = math.sqrt(max(1e-18, ((target_front / 1_000_000.0) / 1.67) ** 2 - 2.5 * (float(values["stray_inductance_uH"]) / 1_000_000.0) * (load_total / 1e12))) / (load_total / 1e12)
    front_stage_theoretical = total_front / stages
    front_stage_selected = _excel_round(front_stage_theoretical / 5.0) * 5.0
    total_tail = (target_tail / 1_000_000.0) / (0.693 * ((generator_cap / 1_000_000.0) + (load_total / 1e12)))
    tail_stage_theoretical = total_tail / stages
    tail_stage_selected = _excel_round(tail_stage_theoretical / 5.0) * 5.0
    physics_front = 1.67 * math.sqrt(((front_stage_selected * stages) * (load_total / 1e12)) ** 2 + 2.5 * (float(values["stray_inductance_uH"]) / 1_000_000.0) * (load_total / 1e12)) * 1_000_000.0
    physics_tail = 0.693 * (tail_stage_selected * stages) * ((generator_cap / 1_000_000.0) + (load_total / 1e12)) * 1_000_000.0
    physics_crest = float(values["test_kV"])
    return {
        **values,
        "E5": target_front,
        "E6": target_tail,
        "E7": load_total,
        "E8": stages,
        "E9": charge_stage,
        "E10": utilization,
        "E11": generator_cap,
        "E12": total_front,
        "E13": front_stage_theoretical,
        "E14": front_stage_selected,
        "E15": total_tail,
        "E16": tail_stage_theoretical,
        "E17": tail_stage_selected,
        "E18": physics_front,
        "E19": physics_tail,
        "E20": physics_crest,
        "B18": front_stage_selected,
        "B19": 1,
        "B20": tail_stage_selected,
        "B21": 1,
    }


def _helper_formula_values(book: Mapping[str, Any], calculator: Mapping[str, Any]) -> dict[str, dict[str, Any]]:
    """Reproduce J:M/O in ML Helper and the three weighted predictions."""

    rows = _sheet_rows(book, "ML Helper")
    distances: list[tuple[int, float]] = []
    raw: dict[int, dict[str, Any]] = {}
    for row_number, row in enumerate(rows, 1):
        if row_number == 1 or row.get("A") is None:
            continue
        impulse = row.get("A")
        distance = (0 if impulse == calculator["impulse_type"] else 1) ** 2
        distance += ((float(row.get("B", 0)) - float(calculator["test_kV"])) / 1500.0) ** 2
        distance += ((float(row.get("C", 0)) - float(calculator["load_pF"])) / 1300.0) ** 2
        distance += ((float(row.get("D", 0)) - float(calculator["divider_pF"])) / 800.0) ** 2
        distance += ((float(row.get("E", 0)) - float(calculator["stray_pF"])) / 350.0) ** 2
        distance += ((float(row.get("F", 0)) - float(calculator["stray_inductance_uH"])) / 30.0) ** 2
        distance += (float(row.get("G", 0)) - float(calculator["efficiency"])) ** 2
        mode_scale = 100.0 if impulse == "Lightning" else 30_000.0
        distance += ((float(row.get("H", 0)) - float(calculator["E14"])) / mode_scale) ** 2
        distance += ((float(row.get("I", 0)) - float(calculator["E17"])) / mode_scale) ** 2
        distances.append((row_number, distance))
        raw[row_number] = row
    distances.sort(key=lambda item: (item[1], item[0]))
    rank_by_row: dict[int, int] = {}
    # Excel RANK(...,1) uses competition ranking: ties share the first rank,
    # leaving gaps after the tie.  Distances are normally unique, but keeping
    # this detail makes the independent reproduction exact for ties too.
    for position, (row_number, distance) in enumerate(distances, 1):
        rank_by_row[row_number] = 1 + sum(1 for _, other in distances if other < distance)
    helper: dict[str, dict[str, Any]] = {}
    weighted = {"N": [], "P": [], "Q": []}
    for row_number, distance in distances:
        row = raw[row_number]
        rank = rank_by_row[row_number]
        use = 1 if rank <= 7 else 0
        weight = 1.0 / (distance + 0.000001) if use else 0.0
        key = str(row_number)
        helper[key] = {"J": distance, "K": rank, "L": use, "M": weight, "O": weight}
        for target, source in (("N", "N"), ("P", "P"), ("Q", "Q")):
            residual = _finite(row.get(source))
            if residual is not None:
                weighted[target].append((residual, weight))
    predictions = {}
    for target, base_key in (("H6", "E18"), ("H7", "E19"), ("H8", "E20")):
        residuals = weighted["N" if target == "H6" else "P" if target == "H7" else "Q"]
        denominator = sum(weight for _, weight in residuals)
        predictions[target] = float(calculator[base_key]) + (sum(value * weight for value, weight in residuals) / denominator if denominator else 0.0)
    return {"helper": helper, "predictions": predictions}


def reproduce_numeric_formulas(path: str | Path) -> dict[str, Any]:
    """Reproduce every cached numeric formula in the original workbook."""

    book = read_workbook(path)
    calculator = _calculator_values(book)
    helper_values = _helper_formula_values(book, calculator)
    expected_calculator = dict((key, calculator[key]) for key in ("E5", "E6", "E7", "E8", "E9", "E10", "E11", "E12", "E13", "E14", "E15", "E16", "E17", "B18", "E18", "E19", "E20", "B19", "B20", "B21"))
    expected_calculator.update({"H6": helper_values["predictions"]["H6"], "H7": helper_values["predictions"]["H7"], "H8": helper_values["predictions"]["H8"]})
    expected_calculator.update({"H9": expected_calculator["H6"] - calculator["E18"], "H10": expected_calculator["H7"] - calculator["E19"], "H11": expected_calculator["H8"] - calculator["E20"]})
    comparisons: list[dict[str, Any]] = []
    formula_count = 0
    numeric_formula_count = 0
    for sheet, cells in book["formulas"].items():
        cached_cells = book["cached"].get(sheet, {})
        for coordinate, item in cells.items():
            formula = item.get("formula")
            if not formula:
                continue
            formula_count += 1
            observed = cached_cells.get(coordinate, {}).get("value")
            observed_numeric = _finite(observed)
            if observed_numeric is None:
                continue
            numeric_formula_count += 1
            expected: float | None = None
            if sheet == "Hybrid Calculator" and coordinate in expected_calculator:
                expected = _finite(expected_calculator[coordinate])
            elif sheet == "ML Helper" and coordinate[0] in {"J", "K", "L", "M", "O"}:
                try:
                    expected = _finite(helper_values["helper"][coordinate[1:]][coordinate[0]])
                except KeyError:
                    expected = None
            absolute_error = None if expected is None else abs(observed_numeric - expected)
            relative_error = None if expected in (None, 0) else absolute_error / abs(expected)
            status = "REPRODUCED" if expected is not None and (math.isclose(observed_numeric, expected, rel_tol=2e-10, abs_tol=2e-10)) else ("UNSUPPORTED_FORMULA" if expected is None else "MISMATCH")
            comparisons.append({"sheet": sheet, "cell": coordinate, "formula": formula, "cached_value": observed_numeric, "reproduced_value": expected, "absolute_error": absolute_error, "relative_error": relative_error, "status": status})
    status_counts: dict[str, int] = defaultdict(int)
    for comparison in comparisons:
        status_counts[comparison["status"]] += 1
    return {"schema_version": COMPARISON_SCHEMA_VERSION, "workbook": {"path": book["path"], "sha256": book["sha256"], "backend": book["backend"], "expected_source_sha256": EXPECTED_SOURCE_SHA256, "source_hash_matches_expected": book["sha256"] == EXPECTED_SOURCE_SHA256}, "formula_count": formula_count, "numeric_formula_count": numeric_formula_count, "status_counts": dict(sorted(status_counts.items())), "calculator": calculator, "comparisons": comparisons}


def _synthetic_rows(book: Mapping[str, Any]) -> list[dict[str, Any]]:
    rows = _sheet_rows(book, "Synthetic Dataset")
    if not rows:
        return []
    headers = rows[0]
    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        record = {header: row.get(column) for column, header in headers.items() if header}
        if record.get("ID") is not None:
            records.append(record)
    return records


def representative_inputs(path: str | Path, *, sample_per_stratum: int = 2, seed: int = 20261010) -> list[dict[str, Any]]:
    """Select original inputs by impulse type, split, and stage band."""

    if type(sample_per_stratum) is not int or sample_per_stratum < 1:
        raise ValueError("sample_per_stratum must be positive")
    book = read_workbook(path)
    records = _synthetic_rows(book)
    buckets: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        stages = int(float(record.get("Stages", 0))) if _finite(record.get("Stages")) is not None else 0
        stage_band = "low" if stages <= 5 else ("mid" if stages <= 10 else "high")
        buckets[(record.get("Impulse_Type"), record.get("Split"), stage_band)].append(record)
    rng = random.Random(seed)
    selected: list[dict[str, Any]] = []
    for key in sorted(buckets, key=lambda value: tuple(str(item) for item in value)):
        values = list(buckets[key])
        rng.shuffle(values)
        selected.extend(copy_record for copy_record in values[:sample_per_stratum])
    return selected


def _configuration_setup(record: Mapping[str, Any]) -> tuple[dict[str, Any], dict[str, Any], str, float]:
    impulse = "LI" if str(record.get("Impulse_Type")) == "Lightning" else "SI"
    stages = int(float(record["Stages"]))
    stage_charge = float(record["Charge_kV_Stage"]) * 1_000.0
    configuration = {
        "impulse_type": impulse,
        "stages": stages,
        "stage_charge_V": stage_charge,
        "front_per_stage_ohm": float(record["Front_R_Stage"]),
        "tail_per_stage_ohm": float(record["Tail_R_Stage"]),
        "polarity": 1,
    }
    setup = {
        "dut_capacitance_F": float(record["Load_C_pF"]) * 1e-12,
        "divider_capacitance_F": float(record["Divider_C_pF"]) * 1e-12,
        "stray_capacitance_F": float(record["Stray_C_pF"]) * 1e-12,
        "loop_inductance_H": float(record["L_uH"]) * 1e-6,
        # The original workbook already sums the visible capacitances; adding
        # the separate 480 pF CPRI basic load would double count it.
        "basic_coverage_assumption": "INCLUDED_IN_OTHER_COMPONENTS",
        "loop_resistance_ohm": 0.0,
        "load_resistance_ohm": None,
        "setup_id": f"EXCEL_ROW_{record.get('ID')}",
        "auxiliary_assumption": "UNCONFIRMED_AUXILIARY_BRANCHES_OMITTED",
    }
    target = float(record["Test_kV"]) * 1_000.0
    return configuration, setup, impulse, target


def _comparison_delta(detailed: Mapping[str, Any], reference: Mapping[str, Any]) -> dict[str, float | None]:
    """Return detailed-model minus one workbook reference for three metrics."""

    result: dict[str, float | None] = {}
    for metric in ("front_us", "tail_us", "crest_kV"):
        actual = _finite(detailed.get(metric))
        expected = _finite(reference.get(metric))
        result[metric] = actual - expected if actual is not None and expected is not None else None
    return result


def _observed_metrics(record: Mapping[str, Any]) -> dict[str, Any]:
    """Normalize the workbook's synthetic observations while retaining columns."""

    front = _finite(record.get("Observed_FrontPeak_us"))
    tail = _finite(record.get("Observed_Tail_us"))
    crest = _finite(record.get("Observed_Crest_kV"))
    return {
        "front_us": front,
        "tail_us": tail,
        "crest_kV": crest,
        "source_columns": {
            "Observed_FrontPeak_us": front,
            "Observed_Tail_us": tail,
            "Observed_Crest_kV": crest,
        },
    }


def _metric_summary(values: Sequence[float]) -> dict[str, Any]:
    finite_values = [float(value) for value in values if _finite(value) is not None]
    if not finite_values:
        return {"count": 0, "mean": None, "median": None, "mae": None}
    return {
        "count": len(finite_values),
        "mean": statistics.fmean(finite_values),
        "median": statistics.median(finite_values),
        "mae": statistics.fmean(abs(value) for value in finite_values),
    }


def _comparison_aggregates(comparisons: Sequence[Mapping[str, Any]], keys: Sequence[str]) -> dict[str, Any]:
    """Aggregate deltas and unsupported reasons by mode/topology dimensions."""

    buckets: dict[tuple[str, ...], list[Mapping[str, Any]]] = defaultdict(list)
    for comparison in comparisons:
        buckets[tuple(str(comparison.get(key)) for key in keys)].append(comparison)
    result: dict[str, Any] = {}
    for bucket_key, rows in sorted(buckets.items()):
        name = "::".join(bucket_key)
        unsupported = [row for row in rows if row.get("detailed_rlc", {}).get("status") != "SIMULATED"]
        unsupported_codes: dict[str, int] = defaultdict(int)
        for row in unsupported:
            code = str((row.get("detailed_rlc", {}).get("error") or {}).get("code", "UNKNOWN"))
            unsupported_codes[code] += 1
        result[name] = {
            "comparison_count": len(rows),
            "simulated_count": len(rows) - len(unsupported),
            "unsupported_count": len(unsupported),
            "unsupported_by_code": dict(sorted(unsupported_codes.items())),
            "delta_detailed_minus_workbook_theoretical": {
                metric: _metric_summary([row.get("delta_detailed_minus_workbook_theoretical", {}).get(metric) for row in rows])
                for metric in ("front_us", "tail_us", "crest_kV")
            },
            "delta_detailed_minus_synthetic_observed": {
                metric: _metric_summary([row.get("delta_detailed_minus_synthetic_observed", {}).get(metric) for row in rows])
                for metric in ("front_us", "tail_us", "crest_kV")
            },
        }
    return result


def compare_detailed_rlc(path: str | Path, *, sample_per_stratum: int = 2, seed: int = 20261010, n_points: int = 1_600) -> dict[str, Any]:
    """Compare representative workbook inputs with 3-uF v3 RLC under both graphs."""

    source = Path(path)
    book = read_workbook(source)
    selected = representative_inputs(source, sample_per_stratum=sample_per_stratum, seed=seed)
    try:
        from powernext_v3.physics import simulate
        from physics_engine.network import PhysicsError
    except ImportError as exc:
        raise RuntimeError("powernext_v3.physics is required for detailed RLC comparison") from exc
    comparisons: list[dict[str, Any]] = []
    for record in selected:
        configuration, setup, impulse, target = _configuration_setup(record)
        workbook_theoretical = {
            "front_us": _finite(record.get("Physics_FrontPeak_us")),
            "tail_us": _finite(record.get("Physics_Tail_us")),
            "crest_kV": _finite(record.get("Physics_Crest_kV")),
        }
        # Future workbook versions may carry measured columns.  Keep those
        # values separate; this release has no laboratory observation rows.
        workbook_observed = _observed_metrics(record)
        for topology in ("GSHUNT_v0", "OSHUNT_v0"):
            request = dict(configuration, topology_id=topology)
            started = datetime.now(timezone.utc)
            try:
                result = simulate(request, setup, domain_id="research_3uf", target_crest_V=target, n_points=n_points, solver="modal")
                metadata = result.metadata
                metrics = metadata.get("metrics", {}) if isinstance(metadata.get("metrics"), Mapping) else {}
                front_value = metrics.get("T1_s") if impulse == "LI" else metrics.get("Tp_s")
                tail_value = metrics.get("T2_s")
                detailed = {"status": "SIMULATED", "front_us": _finite(front_value) * 1e6 if _finite(front_value) is not None else None, "tail_us": _finite(tail_value) * 1e6 if _finite(tail_value) is not None else None, "crest_kV": _finite(metrics.get("crest_magnitude_V")) / 1_000.0 if _finite(metrics.get("crest_magnitude_V")) is not None else None, "numeric_status": metadata.get("numeric_status"), "compliance_status": metrics.get("compliance_status"), "evidence_domain": metadata.get("evidence_domain")}
            except PhysicsError as exc:
                detailed = {"status": "INVALID_OR_UNSUPPORTED", "front_us": None, "tail_us": None, "crest_kV": None, "numeric_status": None, "compliance_status": None, "error": {"type": type(exc).__name__, "code": getattr(exc, "code", None), "message": str(exc)}}
            delta_theoretical = _comparison_delta(detailed, workbook_theoretical)
            delta_observed = _comparison_delta(detailed, workbook_observed)
            comparisons.append({"original_id": record.get("ID"), "original_split": record.get("Split"), "impulse_type": impulse, "domain_id": "research_3uf", "topology_id": topology, "configuration": request, "setup": setup, "workbook_theoretical": workbook_theoretical, "workbook_observed": workbook_observed, "detailed_rlc": detailed, "delta_detailed_minus_workbook_theoretical": delta_theoretical, "delta_detailed_minus_workbook_observed": delta_observed, "delta_detailed_minus_synthetic_observed": delta_observed, "comparison_kind": "MODEL_TO_MODEL", "started_at": started.isoformat()})
    unsupported = [row for row in comparisons if row["detailed_rlc"].get("status") != "SIMULATED"]
    unsupported_by_code: dict[str, int] = defaultdict(int)
    for row in unsupported:
        unsupported_by_code[str((row["detailed_rlc"].get("error") or {}).get("code", "UNKNOWN"))] += 1
    return {"schema_version": COMPARISON_SCHEMA_VERSION, "workbook": {"path": str(source), "sha256": book["sha256"], "source_hash_matches_expected": book["sha256"] == EXPECTED_SOURCE_SHA256}, "domain_id": "research_3uf", "profile_note": "3 uF/stage detailed RLC is a research comparison domain, not CPRI hardware data.", "graph_assumptions": ["GSHUNT_v0", "OSHUNT_v0"], "sample_per_stratum": sample_per_stratum, "selected_input_count": len(selected), "comparison_count": len(comparisons), "unsupported_count": len(unsupported), "unsupported_by_code": dict(sorted(unsupported_by_code.items())), "aggregates": {"by_impulse": _comparison_aggregates(comparisons, ("impulse_type",)), "by_topology": _comparison_aggregates(comparisons, ("topology_id",)), "by_impulse_topology": _comparison_aggregates(comparisons, ("impulse_type", "topology_id"))}, "observed_data_used_for_training": False, "comparisons": comparisons}


def compare_workbook(path: str | Path | None = None, *, sample_per_stratum: int = 2, seed: int = 20261010, n_points: int = 1_600, include_rlc: bool = True) -> dict[str, Any]:
    source = locate_workbook(path) if path is None else Path(path)
    formula_report = reproduce_numeric_formulas(source)
    report: dict[str, Any] = {"schema_version": COMPARISON_SCHEMA_VERSION, "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "source": {"path": str(source), "sha256": _sha256(source), "expected_sha256": EXPECTED_SOURCE_SHA256, "unchanged_check": _sha256(source) == EXPECTED_SOURCE_SHA256}, "formula_reproduction": formula_report, "calculator_method": {"front_heuristic": 1.67, "tail_heuristic": 0.693, "efficiency_input": 0.82, "crest_rule": "Physics crest equals workbook Test_kV target; no observed crest is inferred."}}
    if include_rlc:
        report["detailed_rlc_comparison"] = compare_detailed_rlc(source, sample_per_stratum=sample_per_stratum, seed=seed, n_points=n_points)
    return report


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=None)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sample-per-stratum", type=int, default=2)
    parser.add_argument("--seed", type=int, default=20261010)
    parser.add_argument("--n-points", type=int, default=1_600)
    parser.add_argument("--formula-only", action="store_true")
    args = parser.parse_args(argv)
    report = compare_workbook(args.workbook, sample_per_stratum=args.sample_per_stratum, seed=args.seed, n_points=args.n_points, include_rlc=not args.formula_only)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(__import__("json").dumps(report, indent=2, sort_keys=True, allow_nan=False), encoding="utf-8")
    print(__import__("json").dumps({"output": str(args.output), "formula_status": report["formula_reproduction"]["status_counts"], "source_unchanged": report["source"]["unchanged_check"], "comparison_count": report.get("detailed_rlc_comparison", {}).get("comparison_count", 0)}, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
