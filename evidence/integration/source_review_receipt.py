"""One-time source extraction receipt; requires the document review runtime."""
from pathlib import Path
import hashlib,json
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[2];workspace=root.parent;original=workspace/'sources/original';out=Path(__file__).parent
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
report={'source_precedence':['Original PS','CPRI clarifications and actual IVG sheet','Declared assumptions','Secondary audits'], 'files':[]}
for path in sorted(original.iterdir()):
    report['files'].append(dict(name=path.name,bytes=path.stat().st_size,sha256=sha(path)))
pdf=PdfReader(original/'HV IG Problem Statement.pdf');pages=[p.extract_text() for p in pdf.pages]
(out/'original_problem_statement_all_pages.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{text}' for i,text in enumerate(pages)),encoding='utf-8')
report['problem_statement']={'pages':len(pages),'visual_review':'All four pages including diagrams, page-4 illustrative nameplate and table reviewed in the integration session'}
doc=Document(original/'Parameter values ivg.docx');tables=[[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
(out/'ivg_sheet_all_cells.json').write_text(json.dumps(dict(paragraphs=[p.text for p in doc.paragraphs],tables=tables),ensure_ascii=False,indent=2),encoding='utf-8')
book=load_workbook(original/'Hybrid_Physics_ML_Impulse_Generator_Optimiser.xlsx',read_only=True,data_only=False)
report['workbook']=[]
for sheet in book:
    digest=hashlib.sha256();nonempty=0;formulas=0;rows=0
    for row in sheet.iter_rows(values_only=True):
        rows+=1
        for cell in row:
            if cell is not None:
                nonempty+=1;formulas+=int(isinstance(cell,str) and cell.startswith('='));digest.update(json.dumps(cell,default=str,ensure_ascii=False).encode());digest.update(b'\0')
    report['workbook'].append(dict(sheet=sheet.title,rows_read=rows,columns=sheet.max_column,nonempty_cells=nonempty,formulas=formulas,ordered_nonempty_cell_sha256=digest.hexdigest()))
book.close()
audit=workspace/'archive/audits/PowerNext_Independent_Audit_2026-10-07.pdf';reader=PdfReader(audit)
(out/'independent_audit_all_pages.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{p.extract_text()}' for i,p in enumerate(reader.pages)),encoding='utf-8')
report['independent_audit']=dict(pages=len(reader.pages),sha256=sha(audit),original_location='C:/Users/User/Downloads/RoTrail Web/PowerNext_Independent_Audit_2026-10-07.pdf',preserved_copy=str(audit))
report['model_and_physics_preservation']=[]
baseline=workspace/'archive/PowerNext_CPRI_Revalidation_2026-10-04/release/PowerNext_Track1_Competition_v2'
paths=list((root/'powernext/physics/physics_engine').glob('*.py'))+[root/'powernext/physics/CPRI_EQUIPMENT_PROFILE.json',root/'powernext/physics/CIRCUIT_TOPOLOGY_SPEC.md']+list((root/'powernext/ml/registry/simulation_v2').rglob('model.joblib'))+list((root/'powernext/ml/registry/simulation_v2').rglob('card.json'))
for path in paths:
    previous=baseline/path.relative_to(root);assert sha(path)==sha(previous),path
    report['model_and_physics_preservation'].append(dict(path=path.relative_to(root).as_posix(),sha256=sha(path),identical_to_October4=True))
(out/'source_review_receipt.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(dict(status='PASS',PS_pages=len(pages),IVG_tables=len(tables),workbook_sheets=report['workbook'],audit_pages=len(reader.pages),unchanged_physics_models=len(paths)),indent=2))
