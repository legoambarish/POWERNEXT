"""Read-only workbook ingestion and exact Excel baseline semantics."""
from pathlib import Path
import json
import zipfile
from collections import Counter
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.formula.translate import Translator
from scipy.stats import rankdata
from sklearn.metrics import r2_score
from sklearn.neighbors import NearestNeighbors
from .common import file_hash,write_json,log
from .features import LEGACY_BASE,LEGACY_TARGETS,LEGACY_DIRECT

WORKBOOK_FEATURES=["Test_kV","Load_C_pF","Divider_C_pF","Stray_C_pF","L_uH","Efficiency","Front_R_Stage","Tail_R_Stage"]
RESIDUALS=["Residual_FrontPeak","Residual_Tail","Residual_Crest"]


def ingest(path):
    wb=openpyxl.load_workbook(path,data_only=True,read_only=True)
    rows=list(wb["Synthetic Dataset"].values)
    data=pd.DataFrame(rows[1:],columns=rows[0])
    if list(data.columns)!=["ID","Split","Impulse_Type"]+WORKBOOK_FEATURES[:6]+["Stages","Charge_kV_Stage"]+WORKBOOK_FEATURES[6:]+LEGACY_BASE+LEGACY_TARGETS+RESIDUALS:
        raise ValueError("Workbook schema changed")
    if data.isna().any().any() or data.ID.duplicated().any():
        raise ValueError("Missing/duplicate source records")
    if not np.isfinite(data.select_dtypes("number")).all().all(): raise ValueError("Nonfinite source data")
    return data


def distances(train,query):
    t=train[WORKBOOK_FEATURES].to_numpy(float)
    q=query[WORKBOOK_FEATURES].to_numpy(float)
    # Excel normalizes resistor differences by the TRAINING ROW's type.
    scales=np.tile([1500.,1300.,800.,350.,30.,1.,1.,1.],(len(train),1))
    scales[:,6:]=np.where((train.Impulse_Type=="Lightning").to_numpy()[:,None],100.,30000.)
    d=(((q[:,None,:]-t[None,:,:])/scales[None,:,:])**2).sum(axis=2)
    return d+(query.Impulse_Type.to_numpy()[:,None]!=train.Impulse_Type.to_numpy()[None,:])


def exact_predict(train,query):
    d=distances(train,query)
    # RANK(...,1)<=7 selects every tie at the seventh distance.
    cutoff=np.partition(d,6,axis=1)[:,6]
    w=np.where(d<=cutoff[:,None],1/(d+1e-6),0.)
    return query[LEGACY_BASE].to_numpy(float)+(w@train[RESIDUALS].to_numpy(float))/w.sum(axis=1)[:,None]


def score(y,p):
    e=np.abs(y-p)
    return dict(MAE=e.mean(axis=0).tolist(),RMSE=np.sqrt(((y-p)**2).mean(axis=0)).tolist(),MAPE_pct=(100*e/y).mean(axis=0).tolist(),R2=r2_score(y,p,multioutput="raw_values").tolist())


def audit(path,output):
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    source_hash=file_hash(path)
    data=ingest(path);data.to_csv(output/"legacy.csv",index=False)
    wb=openpyxl.load_workbook(path,data_only=False)
    cache=openpyxl.load_workbook(path,data_only=True)
    inventory=[];all_cells=[]
    for ws in wb:
        formulas=[]
        for row in ws:
            for c in row:
                if c.value is not None:
                    all_cells.append(dict(sheet=ws.title,cell=c.coordinate,value=c.value,cached=cache[ws.title][c.coordinate].value,cell_type=c.data_type))
                    if c.data_type=="f":formulas.append(c.coordinate)
        inventory.append(dict(sheet=ws.title,rows=ws.max_row,columns=ws.max_column,formula_count=len(formulas),state=ws.sheet_state,merged_ranges=[str(x) for x in ws.merged_cells.ranges]))
    write_json(output/"all_cells.json",all_cells)
    write_json(output/"sheet_inventory.json",inventory)
    data_stats=data.groupby("Impulse_Type").describe().T
    data_stats.to_csv(output/"all_column_statistics.csv")
    train=data[data.Split=="Train"]
    helper=cache["ML Helper"]
    expected_helper=train[["Impulse_Type"]+WORKBOOK_FEATURES+RESIDUALS].to_numpy()
    stored_helper=np.array([[helper.cell(r,c).value for c in [1,2,3,4,5,6,7,8,9,14,16,17]] for r in range(2,1402)],dtype=object)
    def multiset(x):return Counter(tuple(r) for r in x)
    helper_matches=multiset(expected_helper)==multiset(stored_helper)
    calc=cache["Hybrid Calculator"]
    query=pd.DataFrame([dict(Impulse_Type=calc["B6"].value,**dict(zip(WORKBOOK_FEATURES,[calc[f"B{r}"].value for r in [7,8,9,10,11,12]]+[calc["E14"].value,calc["E17"].value])),**dict(zip(LEGACY_BASE,[calc[f"E{r}"].value for r in [18,19,20]])))])
    htrain=pd.DataFrame(stored_helper,columns=["Impulse_Type"]+WORKBOOK_FEATURES+RESIDUALS)
    d=distances(htrain,query)[0];rank=rankdata(d,method="min");use=rank<=7;w=np.where(use,1/(d+1e-6),0)
    helper_checks={}
    for col,values in [(10,d),(11,rank),(12,use.astype(int)),(13,w),(15,w)]:
        helper_checks[openpyxl.utils.get_column_letter(col)]=float(np.max(abs(values-np.array([helper.cell(r,col).value for r in range(2,1402)]))))
    formula_mismatches=[]
    for col in ["J","K","L","M","O"]:
        base=wb["ML Helper"][f"{col}2"].value
        for r in range(2,1402):
            coord=f"{col}{r}"
            if Translator(base,origin=f"{col}2").translate_formula(coord)!=wb["ML Helper"][coord].value:
                formula_mismatches.append(coord)
    # Independently rebuild all 29 calculator formulas, including Excel rounding.
    b={r:calc[f"B{r}"].value for r in range(5,16)}
    v={"E5":1.2 if b[6]=="Lightning" else 250.,"E6":50. if b[6]=="Lightning" else 2500.,"E7":sum(b[r] for r in [8,9,10])}
    v["E8"]=int(np.ceil(b[7]/(b[14]*b[12])));v["E9"]=b[7]/(b[12]*v["E8"]);v["E10"]=v["E9"]/b[14];v["E11"]=b[15]/v["E8"]
    cl=v["E7"]*1e-12
    v["E12"]=np.sqrt(max(1e-18,(v["E5"]*1e-6/1.67)**2-2.5*b[11]*1e-6*cl))/cl
    v["E13"]=v["E12"]/v["E8"];v["E14"]=np.floor(v["E13"]/5+.5)*5
    v["E15"]=v["E6"]*1e-6/(.693*(v["E11"]*1e-6+cl));v["E16"]=v["E15"]/v["E8"];v["E17"]=np.floor(v["E16"]/5+.5)*5
    v["E18"]=1.67*np.sqrt((v["E14"]*v["E8"]*cl)**2+2.5*b[11]*1e-6*cl)*1e6
    v["E19"]=.693*v["E17"]*v["E8"]*(v["E11"]*1e-6+cl)*1e6;v["E20"]=b[7]
    for a,value in zip(["H6","H7","H8"],exact_predict(htrain,query)[0]):v[a]=value
    for a,h,e in [("H9","H6","E18"),("H10","H7","E19"),("H11","H8","E20")]:v[a]=v[h]-v[e]
    v.update(B18=np.floor(v["E14"]/5+.5)*5,B19=1,B20=np.floor(v["E17"]/5+.5)*5,B21=1)
    v["H12"]="YES" if v["E8"]<=b[13] and v["E9"]<=b[14] else "NO"
    low,high=(.84,1.56) if b[6]=="Lightning" else (200.,300.)
    v["H13"]="YES" if low<=v["H6"]<=high and v["E6"]*.8<=v["H7"]<=v["E6"]*1.2 and abs(v["H8"]/b[7]-1)<=.03 else "NO"
    v["H14"]="PASS — RECOMMENDED" if v["H12"]==v["H13"]=="YES" else "REVIEW — NOT RECOMMENDED"
    formula_cells={c.coordinate for row in wb["Hybrid Calculator"] for c in row if c.data_type=="f"}
    assert set(v)==formula_cells
    calc_errors={k:(float(abs(value-calc[k].value)) if not isinstance(value,str) else value==calc[k].value) for k,value in v.items()}
    # Baseline reproduction is completed before candidate model selection.
    results=[];predictions=[]
    for split in ["Validation","Hidden Test"]:
        q=data[data.Split==split];pred=exact_predict(train,q)
        for typ in ["ALL","Lightning","Switching"]:
            mask=np.ones(len(q),bool) if typ=="ALL" else (q.Impulse_Type==typ).to_numpy()
            results.append(dict(split=split,type=typ,n=int(mask.sum()),**score(q[LEGACY_TARGETS].to_numpy(float)[mask],pred[mask])))
        predictions.append(pd.concat([q[["ID","Split","Impulse_Type"]].reset_index(drop=True),pd.DataFrame(pred,columns=["pred_front_us","pred_tail_us","pred_crest_kV"])],axis=1))
    held=data[data.Split!="Train"];pooled=score(held[LEGACY_TARGETS].to_numpy(float),exact_predict(train,held))
    hardcoded={key:[cache["Model Validation"].cell(row,col).value for row in [4,5,6]] for key,col in [("MAE",2),("RMSE",3),("MAPE_pct",4),("R2",5)]}
    max_diff=max(float(np.max(abs(np.array(pooled[k])-hardcoded[k]))) for k in pooled)
    pd.concat(predictions).to_csv(output/"exact_baseline_predictions.csv",index=False)
    write_json(output/"exact_baseline_results.json",dict(per_split_type=results,pooled_600=pooled,workbook_displayed=hardcoded,max_pooled_difference=max_diff,calculator=v,calculator_errors=calc_errors,helper_cached_max_errors=helper_checks,helper_formula_mismatches=formula_mismatches,helper_training_match=helper_matches))
    cl=data[["Load_C_pF","Divider_C_pF","Stray_C_pF"]].sum(axis=1)*1e-12
    inferred_cs=(data.Physics_Tail_us*1e-6/(.693*data.Tail_R_Stage*data.Stages)-cl)*data.Stages
    fcalc=1.67*np.sqrt((data.Front_R_Stage*data.Stages*cl)**2+2.5*data.L_uH*1e-6*cl)*1e6
    rules=dict(crest_equals_requested=bool((data.Physics_Crest_kV==data.Test_kV).all()),
               stage_C_implied_F=[float(inferred_cs.min()),float(inferred_cs.max())],
               max_front_formula_error_us=float(abs(fcalc-data.Physics_FrontPeak_us).max()),
               max_charge_formula_error_kV=float(abs(data.Charge_kV_Stage-data.Test_kV/(data.Stages*data.Efficiency)).max()),
               max_stage_selection_error=float(abs(data.Stages-np.ceil(data.Test_kV/(200*data.Efficiency))).max()),
               residual_identity_error=np.max(abs(data[LEGACY_TARGETS].to_numpy()-data[LEGACY_BASE].to_numpy()-data[RESIDUALS].to_numpy()),axis=0).tolist())
    eda={}; stats=[]
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    for typ,frame in data.groupby("Impulse_Type"):
        numeric=frame.select_dtypes("number").drop(columns="ID")
        numeric.corr(method="pearson").to_csv(output/f"{typ}_pearson.csv")
        numeric.corr(method="spearman").to_csv(output/f"{typ}_spearman.csv")
        for col in numeric:
            s=numeric[col]
            stats.append(dict(type=typ,column=col,min=s.min(),p05=s.quantile(.05),median=s.median(),p95=s.quantile(.95),max=s.max(),mean=s.mean(),std=s.std(),unique=s.nunique(),missing=s.isna().sum(),coefficient_variation=float(s.std()/abs(s.mean())) if s.mean()!=0 else None,dominant_fraction=float(s.value_counts(normalize=True).iloc[0])))
        x=frame[LEGACY_DIRECT].to_numpy(float);scale=np.ptp(x,axis=0);scale[scale==0]=1
        nn=NearestNeighbors(n_neighbors=2).fit(x/scale);dist=nn.kneighbors(x/scale)[0][:,1]/np.sqrt(x.shape[1])
        observed=frame[LEGACY_TARGETS].to_numpy()
        # Four bounds unpacked explicitly below.
        flo,fhi,tlo,thi=(.84,1.56,40,60) if typ=="Lightning" else (200,300,1000,4000)
        ok=(observed[:,0]>=flo)&(observed[:,0]<=fhi)&(observed[:,1]>=tlo)&(observed[:,1]<=thi)&(abs(observed[:,2]/frame.Test_kV.to_numpy()-1)<=.03)
        eda[typ]=dict(n=len(frame),split_counts=frame.Split.value_counts().to_dict(),constant_columns=[c for c in numeric if numeric[c].nunique()==1],
                      exact_feature_duplicates=int(frame.duplicated(LEGACY_DIRECT,keep=False).sum()),nearest_normalized_rms_quantiles=np.quantile(dist,[0,.05,.5,.95,1]).tolist(),
                      near_duplicate_count_at_rms_0p02=int((dist<.02).sum()),compliant_count=int(ok.sum()),noncompliant_count=int((~ok).sum()),
                      front_catalog_match_fraction=float(frame.Front_R_Stage.isin([30,46,3700,5000]).mean()),tail_catalog_match_fraction=float(frame.Tail_R_Stage.isin([180] if typ=="Lightning" else [5000]).mean()))
        fig,axes=plt.subplots(2,3,figsize=(12,7))
        for k,(target,resid) in enumerate(zip(LEGACY_TARGETS,RESIDUALS)):
            axes[0,k].hist(frame[target],bins=25,color="#236b8e");axes[0,k].set_title(target.replace("Observed_","Synthetic "))
            axes[1,k].scatter(frame.L_uH,frame[resid],s=5,alpha=.4);axes[1,k].set_xlabel("Inductance (uH)");axes[1,k].set_ylabel(resid)
        fig.suptitle(f"{typ}: target construction and residuals (legacy synthetic)")
        fig.tight_layout();fig.savefig(output/f"{typ}_eda.png",dpi=150);plt.close(fig)
    pd.DataFrame(stats).to_csv(output/"per_type_feature_statistics.csv",index=False)
    summary=dict(domain="LEGACY_SYNTHETIC_3uF",source_sha256=source_hash,rows=len(data),rules=rules,eda=eda,formula_count=sum(s["formula_count"] for s in inventory),source_unmodified=file_hash(path)==source_hash,baseline_reproduced=max_diff<1e-10 and helper_matches and not formula_mismatches)
    write_json(output/"audit_summary.json",summary)
    assert summary["baseline_reproduced"] and summary["source_unmodified"]
    log(f"Legacy audit complete: {len(data)} records; displayed metric max difference {max_diff:.3g}")
    return summary
